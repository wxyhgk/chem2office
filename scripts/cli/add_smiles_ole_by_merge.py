#!/usr/bin/env python3
"""Add a ChemDraw OLE structure column to an existing data workbook.

For each selected sheet, reads SMILES from column A (rows 2..N), inserts a new
column B titled "Struct", and fills it with editable ChemDraw OLE objects.
Existing columns from B onward shift one column to the right.

Built on the retain_cdx public API: openpyxl performs the plain-data pre-pass
(read SMILES, insert the column, set widths/heights), then ``ExcelCDXWriter``
inserts the OLE objects.

Caveat: the input workbook must not already contain OLE objects or drawings —
the openpyxl pre-pass would drop them. (The legacy XML-merge implementation
had the same practical assumption: it rebuilt the sheet's drawing parts.)

Usage::

    python3 -m scripts.cli.add_smiles_ole_by_merge tests/data/input/data1.xlsx \
        -o tests/output/data1_ole_merge.xlsx --sheets P型 N型 --limit 3
"""

from __future__ import annotations

import argparse
from pathlib import Path

from retain_cdx import (
    DisplayOptions,
    ExcelCDXWriter,
    Placement,
    PreviewOptions,
    StructureSource,
)

DEFAULT_SHEETS = ("P型", "N型")

def write_row_slice(
    input_xlsx: Path,
    slice_xlsx: Path,
    *,
    sheet: str,
    start: int,
    count: int,
) -> int:
    """Copy the header plus data rows ``[start, start+count)`` to a new workbook.

    Used by ``--split``: a few hundred OLE objects per file keeps Excel
    responsive on modest hardware, where thousands on one sheet do not.
    Cell formatting is not carried over — these are plain data dumps.
    """
    from openpyxl import Workbook, load_workbook

    source = load_workbook(input_xlsx, read_only=True)
    if sheet not in source.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet}")
    rows = list(source[sheet].iter_rows(values_only=True))
    source.close()
    if not rows:
        raise SystemExit(f"No rows in {input_xlsx}")

    header, data = rows[0], rows[1:]
    chunk = data[start : start + count]

    book = Workbook()
    target = book.active
    target.title = sheet
    target.append(list(header))
    for row in chunk:
        target.append(list(row))
    book.save(slice_xlsx)
    return len(chunk)


def count_data_rows(input_xlsx: Path, sheet: str) -> int:
    from openpyxl import load_workbook

    book = load_workbook(input_xlsx, read_only=True)
    if sheet not in book.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet}")
    total = max(0, book[sheet].max_row - 1)
    book.close()
    return total


def prepare_workbook(
    input_xlsx: Path,
    base_xlsx: Path,
    *,
    sheets: list[str],
    limit: int,
    row_height: float,
    struct_col_width: float,
    keep_smiles: bool = True,
) -> tuple[dict[str, list[str]], str]:
    """Insert the Struct column per sheet.

    Returns ``({sheet: [smiles, ...]}, struct_column_letter)``. With
    *keep_smiles* the SMILES text stays in column A and structures land in B;
    otherwise column A is dropped and structures take its place.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("openpyxl is required for the data pre-pass.") from exc

    workbook = load_workbook(input_xlsx)
    smiles_by_sheet: dict[str, list[str]] = {}
    for name in sheets:
        if name not in workbook.sheetnames:
            raise SystemExit(f"Sheet not found: {name}")
        worksheet = workbook[name]

        data_rows = min(limit, max(0, worksheet.max_row - 1))
        smiles: list[str] = []
        for row in range(2, data_rows + 2):
            value = worksheet.cell(row=row, column=1).value
            text = str(value).strip() if value is not None else ""
            if not text:
                raise SystemExit(f"{name}!A{row}: empty SMILES cell")
            smiles.append(text)
        smiles_by_sheet[name] = smiles

        if keep_smiles:
            struct_col_index = 2
            worksheet.insert_cols(2)
            worksheet.column_dimensions["A"].width = 55
        else:
            # Drop the SMILES text; structures take column A.
            struct_col_index = 1
            worksheet.delete_cols(1)
            worksheet.insert_cols(1)
        struct_letter = get_column_letter(struct_col_index)
        worksheet.cell(row=1, column=struct_col_index, value="Struct")
        worksheet.column_dimensions[struct_letter].width = struct_col_width
        worksheet.column_dimensions[get_column_letter(struct_col_index + 1)].width = 18
        for row in range(2, data_rows + 2):
            worksheet.row_dimensions[row].height = row_height

    workbook.save(base_xlsx)
    return smiles_by_sheet, get_column_letter(1 if not keep_smiles else 2)


def convert_one(input_xlsx: Path, output: Path, args, *, limit: int) -> None:
    """Add the structure column for one (input, output) pair."""
    base = output.with_name(f".{output.stem}_base.tmp.xlsx")
    smiles_by_sheet, struct_col = prepare_workbook(
        input_xlsx,
        base,
        sheets=args.sheets,
        limit=limit,
        row_height=args.row_height,
        struct_col_width=args.struct_col_width,
        keep_smiles=not args.no_smiles_column,
    )
    start_cell = f"{struct_col}2"

    display = DisplayOptions(
        width="36mm",
        height="26mm",
        row_height=args.row_height,
        col_width=args.struct_col_width,
        preview=PreviewOptions(
            format=args.preview_format,
            width=args.preview_width,
            height=args.preview_height,
            thick_lines=not args.preview_thin_lines,
            no_margin=not args.preview_keep_margin,
        ),
    )

    try:
        with ExcelCDXWriter(base) as writer:
            for name, smiles in smiles_by_sheet.items():
                if not smiles:
                    continue
                sources = []
                for idx, text in enumerate(smiles, start=2):
                    try:
                        sources.append(
                            StructureSource.from_smiles(
                                text,
                                name=f"{name}_row{idx}",
                                bond_length=args.bond_length,
                                line_width=args.line_width,
                                bold_width=args.bold_width,
                                label_size=args.label_size,
                            )
                        )
                    except Exception as exc:
                        raise SystemExit(f"{name}!A{idx}: {exc}") from exc
                insert = (
                    writer.insert_pictures_many
                    if args.picture_mode
                    else writer.insert_cdx_many
                )
                insert(
                    sources,
                    sheet=name,
                    start_cell=start_cell,
                    placement=Placement(start_cell, offset_x="2mm", offset_y="2mm"),
                    display=display,
                )
            writer.save(output)
    finally:
        base.unlink(missing_ok=True)

    print(f"Wrote {output}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Add a ChemDraw OLE structure column to a data workbook, sheet by sheet."
    )
    parser.add_argument("input", nargs="?", default="tests/data/input/data1.xlsx")
    parser.add_argument("-o", "--output", default="tests/output/data1_ole_merge.xlsx")
    parser.add_argument("--sheets", nargs="+", default=list(DEFAULT_SHEETS))
    parser.add_argument("--limit", type=int, default=3)
    parser.add_argument(
        "--all",
        action="store_true",
        help="Process all data rows in each selected sheet.",
    )
    parser.add_argument("--row-height", type=float, default=110)
    parser.add_argument("--struct-col-width", type=float, default=24)
    parser.add_argument("--bond-length", type=float, default=24)
    parser.add_argument("--line-width", type=float, default=1.2)
    parser.add_argument("--bold-width", type=float, default=4)
    parser.add_argument("--label-size", type=float, default=16)
    parser.add_argument(
        "--preview-format",
        choices=("png", "emf", "auto"),
        default="png",
        help="Preview image format; emf is vector (sharp, small) and needs Inkscape",
    )
    parser.add_argument("--preview-width", type=int, default=900)
    parser.add_argument("--preview-height", type=int, default=650)
    parser.add_argument("--preview-thin-lines", action="store_true")
    parser.add_argument("--preview-keep-margin", action="store_true")
    parser.add_argument(
        "--picture-mode",
        action="store_true",
        help=(
            "Insert static pictures instead of editable ChemDraw OLE objects. "
            "Thousands of OLE objects bog Excel down even on capable machines; "
            "pictures carry no OLE container or VML shape, so one sheet holds "
            "thousands of structures comfortably. Trade-off: no double-click editing."
        ),
    )
    parser.add_argument(
        "--no-smiles-column",
        action="store_true",
        help=(
            "Drop the SMILES text column; structures take column A. Note this is "
            "a layout choice, not a performance one — the SMILES text costs well "
            "under 1%% of the file."
        ),
    )
    parser.add_argument(
        "--split",
        type=int,
        default=None,
        metavar="N",
        help=(
            "Write one file per N structures (e.g. 250). Thousands of OLE objects "
            "on a single sheet bog Excel down on modest hardware; splitting keeps "
            "every file responsive and fully editable. Single sheet only."
        ),
    )
    args = parser.parse_args()

    input_xlsx = Path(args.input)
    if not input_xlsx.exists():
        raise SystemExit(f"Input not found: {input_xlsx}")
    output = Path(args.output)
    if output.parent != Path(""):
        output.parent.mkdir(parents=True, exist_ok=True)

    limit = 10**9 if args.all else args.limit

    if args.split is None:
        convert_one(input_xlsx, output, args, limit=limit)
        return

    if args.split < 1:
        raise SystemExit("--split needs a positive row count")
    if len(args.sheets) != 1:
        raise SystemExit("--split supports a single --sheets value")

    sheet = args.sheets[0]
    total = min(limit, count_data_rows(input_xlsx, sheet))
    if total == 0:
        raise SystemExit(f"No data rows in {input_xlsx}")
    parts = (total + args.split - 1) // args.split
    width = max(2, len(str(parts)))
    print(f"{input_xlsx.name}: {total} structures → {parts} files of ≤{args.split}")

    for index in range(parts):
        start = index * args.split
        part_out = output.with_name(f"{output.stem}_{index + 1:0{width}d}{output.suffix}")
        slice_path = output.with_name(f".{part_out.stem}_slice.tmp.xlsx")
        try:
            rows = write_row_slice(
                input_xlsx,
                slice_path,
                sheet=sheet,
                start=start,
                count=min(args.split, total - start),
            )
            print(f"  [{index + 1}/{parts}] rows {start + 2}-{start + rows + 1}")
            convert_one(slice_path, part_out, args, limit=rows)
        finally:
            slice_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
