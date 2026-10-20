from __future__ import annotations

import tempfile
from collections.abc import Iterable
from pathlib import Path

from retain_cdx.ole import word_streams, write_cdx_to_ole
from retain_cdx.render.preview import (
    preview_extension,
    render_preview,
    render_previews_at_natural_scale,
)
from retain_cdx.resources import bundled_template_docx
from retain_cdx.structure import Compound, CompoundTable, StructureSource

from retain_cdx.office._package import extract_ooxml, pack_ooxml
from retain_cdx.office._units import parse_emu
from retain_cdx.office.options import PreviewOptions

from ._wml import (
    CELL_MARGIN_TWIPS,
    ObjectSize,
    OleObjectTemplate,
    TableCell,
    WordDocument,
    build_compound_table,
    check_alignment,
    make_blank,
)

_DEFAULT_WIDTH_EMU = parse_emu("4cm")

Length = int | float | str


class WordCDXWriter:
    """Put editable ChemDraw CDX OLE objects into an existing Word document.

    Scope matches the PowerPoint writer: *this class puts molecules into a
    document*. Writing the document itself — headings, tables, text — belongs
    to a library such as ``python-docx``. The convenient hand-off is a text
    placeholder::

        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "咖啡因"
        table.cell(0, 1).text = "{{caffeine}}"
        doc.save("report.docx")

        with WordCDXWriter("report.docx") as writer:
            writer.insert_cdx(caffeine, placeholder="{{caffeine}}", height="2cm")
            writer.save("report_with_structures.docx")

    Objects are inline: they flow with the text like a character, so there is
    no position to give — only *where in the text* (a placeholder, the end of a
    paragraph, or a new paragraph) and a size.
    """

    def __init__(self, input_docx: str | Path | None = None):
        """Open *input_docx*, or start an empty A4 document when it is ``None``.

        The empty document is the bundled template with its content removed:
        Word's own styles, theme and page setup, nothing else.
        """
        self._tmp = tempfile.TemporaryDirectory(prefix="retain_cdx_docx_")
        self.package_dir = Path(self._tmp.name) / "package"
        self.package_dir.mkdir()
        if input_docx is None:
            with bundled_template_docx() as template:
                extract_ooxml(template, self.package_dir)
            make_blank(self.package_dir)
        else:
            extract_ooxml(Path(input_docx), self.package_dir)
        self._template = self._load_object_template()

    def _load_object_template(self) -> OleObjectTemplate:
        with tempfile.TemporaryDirectory(prefix="retain_cdx_docx_tpl_") as tmp:
            template_dir = Path(tmp)
            with bundled_template_docx() as template:
                extract_ooxml(template, template_dir)
            return OleObjectTemplate.from_package(template_dir)

    def close(self) -> None:
        self._tmp.cleanup()

    def __enter__(self) -> "WordCDXWriter":
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        self.close()

    @property
    def paragraph_count(self) -> int:
        """Top-level paragraphs (what ``paragraph=`` indexes, 0-based)."""
        return len(WordDocument(self.package_dir).body_paragraphs())

    def save(self, output_docx: str | Path) -> None:
        pack_ooxml(self.package_dir, Path(output_docx))

    def insert_cdx(
        self,
        source: StructureSource | str | Path,
        *,
        placeholder: str | None = None,
        paragraph: int | None = None,
        width: Length | None = None,
        height: Length | None = None,
        align: str | None = None,
        preview: PreviewOptions | None = None,
    ) -> None:
        """Insert one structure.

        Where it goes:

        * *placeholder* — replaces the first occurrence of that text (in body
          paragraphs or table cells);
        * *paragraph* — appended to the end of that top-level paragraph
          (0-based, the same order as python-docx's ``document.paragraphs``);
        * neither — a new paragraph at the end of the document, aligned by
          *align* (``"left"``, ``"center"`` or ``"right"``).

        *width* / *height* accept EMU or ``"cm"``/``"mm"``/``"in"``/``"pt"``/
        ``"px"`` strings. Give one and the other follows the molecule's aspect
        ratio; give both and the molecule is fitted inside that box without
        stretching. Default: 4 cm wide.
        """
        self.insert_cdx_many(
            [source],
            placeholders=None if placeholder is None else [placeholder],
            paragraph=paragraph,
            width=width,
            height=height,
            align=align,
            preview=preview,
        )

    def insert_cdx_many(
        self,
        sources: Iterable[StructureSource | str | Path],
        *,
        placeholders: list[str] | None = None,
        paragraph: int | None = None,
        width: Length | None = None,
        height: Length | None = None,
        align: str | None = None,
        preview: PreviewOptions | None = None,
    ) -> None:
        """Insert several structures.

        With *placeholders* (one per structure) each replaces its own text;
        with *paragraph* they are appended side by side to that paragraph;
        otherwise each gets a new paragraph at the end of the document.
        """
        structures = [
            item if isinstance(item, StructureSource) else StructureSource.from_cdx(item)
            for item in sources
        ]
        if not structures:
            return
        if placeholders is not None and paragraph is not None:
            raise ValueError("Give placeholders or paragraph, not both")
        if placeholders is not None and len(placeholders) != len(structures):
            raise ValueError(
                f"placeholders has {len(placeholders)} entries for "
                f"{len(structures)} structures"
            )
        if align is not None and (placeholders is not None or paragraph is not None):
            raise ValueError("align only applies to new paragraphs")

        preview_options = preview or PreviewOptions()
        extension = preview_extension(preview_options.format)
        cache = preview_options.presentation_cache
        if cache and extension != "emf":
            raise ValueError(
                "presentation_cache needs a vector preview; pass PreviewOptions(format='emf')"
            )

        # Check every location before rendering anything: a failure halfway
        # would leave the parts already written as orphans in the package.
        check_alignment(align)
        document = WordDocument(self.package_dir)
        if paragraph is not None:
            count = len(document.body_paragraphs())
            if not 0 <= paragraph < count:
                raise ValueError(
                    f"Paragraph {paragraph} out of range (document has {count})"
                )
        if placeholders is not None:
            for placeholder in dict.fromkeys(placeholders):
                wanted = placeholders.count(placeholder)
                found = document.count_placeholder(placeholder)
                if found < wanted:
                    raise ValueError(
                        f"Placeholder {placeholder!r} not found in the document"
                        if found == 0
                        else f"Placeholder {placeholder!r} occurs {found} times, "
                        f"{wanted} structures given for it"
                    )

        with tempfile.TemporaryDirectory(prefix="retain_cdx_docx_parts_") as tmp:
            work = Path(tmp)
            for index, structure in enumerate(structures):
                box_w, box_h = _resolve_size(structure, width, height)
                cdx_path = work / f"structure{index + 1}.cdx"
                structure.write_cdx(cdx_path)

                image_path = work / f"image{index + 1}.{extension}"
                render_preview(
                    cdx_path,
                    image_path,
                    format=preview_options.format,
                    width=preview_options.width,
                    height=preview_options.height,
                    thick_lines=preview_options.thick_lines,
                    no_margin=preview_options.no_margin,
                    mono=preview_options.mono,
                    backend=preview_options.backend,
                    style=preview_options.style,
                    target_aspect=box_w / box_h,
                )
                preview_bytes = image_path.read_bytes()

                # Word 的渲染缓存就是预览 EMF 本身（\x03EPRINT），不用另跑管线。
                eprint = preview_bytes if extension == "emf" and cache is not False else None
                container_path = work / f"oleObject{index + 1}.bin"
                write_cdx_to_ole(
                    container_path, cdx_path, extra_streams=word_streams(eprint)
                )

                run = document.build_object_run(
                    template=self._template,
                    container=container_path.read_bytes(),
                    preview_bytes=preview_bytes,
                    preview_ext=extension,
                    size=ObjectSize.from_emu(box_w, box_h),
                )
                if placeholders is not None:
                    document.replace_placeholder(placeholders[index], run)
                elif paragraph is not None:
                    document.append_to_paragraph(paragraph, run)
                else:
                    document.append_paragraph(run, align=align)

        document.save()


    def insert_compound_table(
        self,
        compounds: CompoundTable | Iterable[Compound],
        *,
        columns: int | None = None,
        scale: float | None = None,
        placeholder: str | None = None,
        vertical_align: str = "bottom",
        preview: PreviewOptions | None = None,
    ) -> float:
        """Insert a patent-style compound table; returns the scale used.

        Each group of *columns* compounds takes two rows: the structures, then
        their numbers centred underneath. Every structure is shown at its
        natural size (ACS 1996, bond length 14.4 pt) times **one** shared
        *scale*, so bond lengths agree across the whole table — fitting each
        molecule into a fixed box would not.

        *scale* defaults to the largest value at which the widest structure
        still fits its column. *columns* defaults to the table's own setting
        (3 when given a plain list). With *placeholder* the table replaces the
        paragraph holding exactly that text; otherwise it is appended.
        *vertical_align* (``"bottom"``, ``"center"``, ``"top"``) places
        structures of different heights within their row.
        """
        if isinstance(compounds, CompoundTable):
            columns = columns or compounds.columns
            items = list(compounds.compounds)
        else:
            items = list(compounds)
        columns = columns or 3
        if not items:
            raise ValueError("No compounds to put in the table")
        if columns < 1:
            raise ValueError(f"columns must be at least 1, got {columns}")
        if scale is not None and scale <= 0:
            raise ValueError(f"scale must be positive, got {scale}")

        preview_options = preview or PreviewOptions()
        cache = preview_options.presentation_cache
        extension = preview_extension(preview_options.format)
        if cache and extension != "emf":
            raise ValueError(
                "presentation_cache needs a vector preview; pass PreviewOptions(format='emf')"
            )

        document = WordDocument(self.package_dir)
        if placeholder is not None:
            document.find_placeholder_paragraph(placeholder)  # fail before rendering
        column_twips = document.text_width_twips() // columns
        usable_pt = (column_twips - 2 * CELL_MARGIN_TWIPS) / 20

        with tempfile.TemporaryDirectory(prefix="retain_cdx_docx_table_") as tmp:
            work = Path(tmp)
            jobs = []
            for index, compound in enumerate(items):
                cdx_path = compound.structure.write_cdx(work / f"structure{index + 1}.cdx")
                jobs.append((cdx_path, work / f"image{index + 1}.{extension}"))
            resolved, sizes = render_previews_at_natural_scale(
                jobs,
                format=preview_options.format,
                mono=preview_options.mono,
                style=preview_options.style,
                backend=preview_options.backend,
            )
            if scale is None:
                scale = usable_pt / max(width for width, _ in sizes)

            cells = []
            for compound, (cdx_path, image_path), (width, height) in zip(items, jobs, sizes):
                preview_bytes = image_path.read_bytes()
                eprint = preview_bytes if resolved == "emf" and cache is not False else None
                container_path = cdx_path.with_suffix(".bin")
                write_cdx_to_ole(container_path, cdx_path, extra_streams=word_streams(eprint))
                run = document.build_object_run(
                    template=self._template,
                    container=container_path.read_bytes(),
                    preview_bytes=preview_bytes,
                    preview_ext=resolved,
                    size=ObjectSize(width * scale, height * scale),
                )
                cells.append(TableCell(run=run, label=compound.id))

        table = build_compound_table(
            cells, columns=columns, column_twips=column_twips, vertical_align=vertical_align
        )
        document.insert_block(table, placeholder=placeholder)
        document.save()
        return scale


def _resolve_size(
    structure: StructureSource, width: Length | None, height: Length | None
) -> tuple[int, int]:
    """(width, height) in EMU, always in the molecule's own aspect ratio.

    A box that disagrees with the CDX gets the molecule stretched as soon as
    anything re-renders it from CONTENTS — the bug the Excel side shipped with
    until preserve_aspect became the default.
    """
    aspect = structure.content_aspect or 1.0
    width_emu = None if width is None else parse_emu(width)
    height_emu = None if height is None else parse_emu(height)
    if width_emu is None and height_emu is None:
        width_emu = _DEFAULT_WIDTH_EMU
    if height_emu is None:
        return width_emu, max(1, round(width_emu / aspect))
    if width_emu is None:
        return max(1, round(height_emu * aspect)), height_emu
    # Both given: the largest box with the molecule's aspect that fits.
    if width_emu / height_emu > aspect:
        return max(1, round(height_emu * aspect)), height_emu
    return width_emu, max(1, round(width_emu / aspect))


__all__ = ["WordCDXWriter"]
