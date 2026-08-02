from __future__ import annotations

import tempfile
from collections.abc import Iterable
from pathlib import Path

from retain_cdx.ole import (
    MAX_EMBEDDED_EMF,
    build_enhmetafile_presentation,
    emu_to_himetric,
    write_cdx_to_ole,
)
from retain_cdx.render.preview import preview_extension, render_preview
from retain_cdx.resources import bundled_template_pptx
from retain_cdx.structure import StructureSource

from retain_cdx.office._package import extract_ooxml, pack_ooxml
from retain_cdx.office._units import parse_emu

from ._pml import (
    FramePlacement,
    OleFrameTemplate,
    append_ole_frame,
    slide_paths,
    slide_size,
)
from .options import PreviewOptions, PowerPointGrid, PowerPointPlacement

_DEFAULT_WIDTH_EMU = parse_emu("6cm")


class PowerPointCDXWriter:
    """Place editable ChemDraw CDX OLE objects onto slides of an existing deck.

    Scope is deliberately narrow: *this class puts molecules on a slide at a
    given position*. Building the presentation itself — slides, layouts, titles,
    tables — belongs to a presentation library such as ``python-pptx``; hand the
    resulting ``.pptx`` here afterwards::

        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "候选分子"
        prs.save("deck.pptx")

        with PowerPointCDXWriter("deck.pptx") as writer:
            writer.insert_cdx(source, slide=1,
                              placement=PowerPointPlacement(x="2cm", y="5cm", width="5cm"))
            writer.save("deck_with_structures.pptx")

    Either order works: python-pptx round-trips the OLE objects written here
    without losing them (unlike openpyxl on the Excel side).

    Positions are absolute EMU boxes — PresentationML has no cell grid, so none
    of the Excel-side column-width/anchor machinery applies here.
    """

    def __init__(self, input_pptx: str | Path):
        self._tmp = tempfile.TemporaryDirectory(prefix="retain_cdx_pptx_")
        self.package_dir = Path(self._tmp.name) / "package"
        self.package_dir.mkdir()
        extract_ooxml(Path(input_pptx), self.package_dir)
        self._template = self._load_frame_template()

    def _load_frame_template(self) -> OleFrameTemplate:
        with tempfile.TemporaryDirectory(prefix="retain_cdx_pptx_tpl_") as tmp:
            template_dir = Path(tmp)
            with bundled_template_pptx() as template:
                extract_ooxml(template, template_dir)
            return OleFrameTemplate.from_package(template_dir)

    def close(self) -> None:
        self._tmp.cleanup()

    def __enter__(self) -> "PowerPointCDXWriter":
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        self.close()

    @property
    def slide_count(self) -> int:
        return len(slide_paths(self.package_dir))

    @property
    def slide_size(self) -> tuple[int, int]:
        """(width, height) of the slides, in EMU."""
        return slide_size(self.package_dir)

    def save(self, output_pptx: str | Path) -> None:
        pack_ooxml(self.package_dir, Path(output_pptx))

    def insert_cdx(
        self,
        source: StructureSource | str | Path,
        *,
        slide: int = 1,
        placement: PowerPointPlacement | None = None,
        preview: PreviewOptions | None = None,
        name: str | None = None,
    ) -> None:
        """Insert one structure on *slide* (1-based).

        Without *placement* the structure is centred on the slide at a default
        6 cm width, keeping the molecule's aspect ratio.
        """
        self.insert_cdx_many(
            [source], slide=slide, placements=[placement], preview=preview, names=None
            if name is None
            else [name],
        )

    def insert_cdx_many(
        self,
        sources: Iterable[StructureSource | str | Path],
        *,
        slide: int = 1,
        grid: PowerPointGrid | None = None,
        placements: list[PowerPointPlacement | None] | None = None,
        preview: PreviewOptions | None = None,
        names: list[str] | None = None,
    ) -> None:
        """Insert several structures on one slide.

        Give *grid* to tile them, or *placements* for explicit boxes (one entry
        per structure; ``None`` entries fall back to centred defaults).
        """
        structures = [
            item if isinstance(item, StructureSource) else StructureSource.from_cdx(item)
            for item in sources
        ]
        if not structures:
            return

        slide_files = slide_paths(self.package_dir)
        if not 1 <= slide <= len(slide_files):
            raise ValueError(
                f"Slide {slide} out of range (presentation has {len(slide_files)})"
            )
        slide_path = slide_files[slide - 1]

        preview_options = preview or PreviewOptions(format="emf")
        extension = preview_extension(preview_options.format)

        with tempfile.TemporaryDirectory(prefix="retain_cdx_pptx_parts_") as tmp:
            work = Path(tmp)
            for index, structure in enumerate(structures):
                box = self._resolve_box(
                    structure,
                    index=index,
                    grid=grid,
                    placement=placements[index]
                    if placements and index < len(placements)
                    else None,
                )
                cdx_path = work / f"structure{index + 1}.cdx"
                structure.write_cdx(cdx_path)

                image_path = work / f"image{index + 1}.{extension}"
                render_preview(
                    cdx_path,
                    image_path,
                    format=preview_options.format,
                    width=preview_options.width,
                    height=preview_options.height,
                    thick_lines=preview_options.thick_lines,
                    no_margin=preview_options.no_margin,
                    mono=preview_options.mono,
                    backend=preview_options.backend,
                    style=preview_options.style,
                    target_aspect=box.width / box.height if box.height else None,
                    # 缓存渲染要把 EMF 藏进一份 WMF 里，两个都从同一份 SVG 导出。
                    companion_wmf_dir=work if extension == "emf" else None,
                )

                presentation = (
                    None
                    if preview_options.presentation_cache is False
                    else self._build_presentation(image_path, box)
                )

                container_path = work / f"oleObject{index + 1}.bin"
                write_cdx_to_ole(
                    container_path,
                    cdx_path,
                    template_ole=self._template.container,
                    stream_name="CONTENTS",
                    presentation=presentation,
                )

                append_ole_frame(
                    self.package_dir,
                    slide_path=slide_path,
                    template=self._template,
                    container=container_path.read_bytes(),
                    preview_bytes=image_path.read_bytes(),
                    preview_ext=extension,
                    placement=box,
                    name=(names[index] if names and index < len(names) else None)
                    or structure.name
                    or f"Structure {index + 1}",
                )

    def _build_presentation(
        self, image_path: Path, box: FramePlacement
    ) -> bytes | None:
        """缓存渲染（`\x02OlePres000`）——没有它，复制出去只会得到一张图片。

        只有 EMF 预览才有；EMF 必须整份放进一条 WMFC 注释记录，超过
        `MAX_EMBEDDED_EMF` 时宁可不挂（复制退化好过让 PowerPoint 崩）。
        """
        if image_path.suffix != ".emf":
            return None
        emf = image_path.read_bytes()
        if len(emf) > MAX_EMBEDDED_EMF:
            return None
        wmf_path = image_path.with_suffix(".wmf")  # work 目录里，不进包
        if not wmf_path.exists():
            return None
        return build_enhmetafile_presentation(
            emf,
            wmf_path.read_bytes(),
            width_himetric=emu_to_himetric(box.width),
            height_himetric=emu_to_himetric(box.height),
        )

    def _resolve_box(
        self,
        structure: StructureSource,
        *,
        index: int,
        grid: PowerPointGrid | None,
        placement: PowerPointPlacement | None,
    ) -> FramePlacement:
        aspect = structure.content_aspect or 1.0
        if grid is not None:
            x, y, cell_w, cell_h = grid.box(index)
            width, height = _fit(cell_w, cell_h, aspect)
            # Centre the molecule inside its cell.
            return FramePlacement(
                x + (cell_w - width) // 2, y + (cell_h - height) // 2, width, height
            )

        if placement is None:
            width = _DEFAULT_WIDTH_EMU
            height = max(1, round(width / aspect))
            slide_w, slide_h = self.slide_size
            return FramePlacement(
                (slide_w - width) // 2, (slide_h - height) // 2, width, height
            )

        width, height = placement.width_emu, placement.height_emu
        if width is None and height is None:
            width = _DEFAULT_WIDTH_EMU
        if width is None:
            width = max(1, round(height * aspect))
        elif height is None:
            height = max(1, round(width / aspect))
        return FramePlacement(placement.x_emu, placement.y_emu, width, height)


def _fit(width: int, height: int, aspect: float) -> tuple[int, int]:
    """Largest box with *aspect* that fits inside width×height."""
    if width / height > aspect:
        return max(1, round(height * aspect)), height
    return width, max(1, round(width / aspect))


__all__ = ["PowerPointCDXWriter"]
